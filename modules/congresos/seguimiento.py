"""Seguimiento de invitados del Congreso Internacional de octubre de 2026."""

from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from typing import Any
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from database.db import (
    actualizar_confirmacion_proyeccion,
    actualizar_institucion_proyeccion,
    actualizar_invitado_congreso,
    actualizar_responsable_congreso,
    contar_invitados_congreso,
    crear_institucion_proyeccion,
    crear_responsable_congreso,
    get_connection,
    importar_datos_congreso,
    listar_historial_congreso,
    listar_invitados_congreso,
    listar_proyeccion_estudiantes,
    listar_responsables_congreso,
    precargar_proyeccion_estudiantes,
    sincronizar_documentos_congreso,
)


COLOR_AZUL = "#1A3A5C"
# Contrato de compatibilidad con app.py. Se incrementa cuando cambia la
# sincronización que transforma registros ya existentes.
CONGRESO_SYNC_VERSION = 7
ZONA_HORARIA_ECUADOR = ZoneInfo("America/Guayaquil")
ESTADOS = ["Pendiente", "Sí", "No"]
OFICINAS = {
    "guayaquil": "Guayaquil",
    "manabi": "Portoviejo",
    "loja": "Loja",
    "cuenca": "Cuenca",
}
MAPEO_RESPONSABLE_OFICINA = {
    "otap": "manabi",
    "otac": "cuenca",
    "otal": "loja",
    "milka": "guayaquil",
    "carlos g": "guayaquil",
    "despacho": "guayaquil",
    "intendente regional": "guayaquil",
}
MAPEO_RESPONSABLE_DATOS = {
    "milka": ("Ing. Milka Nazareno", "0981768984", "milka.nazareno@sce.gob.ec"),
    "carlos g": ("Ab. Carlos García", "0996797882", "carlos.garcia@sce.gob.ec"),
    "despacho": ("Ab. Carlos García", "0996797882", "carlos.garcia@sce.gob.ec"),
    "intendente regional": (
        "Ing. Milka Nazareno",
        "0981768984",
        "milka.nazareno@sce.gob.ec",
    ),
    "otap": ("Ab. Cristian Romero", "0993345666", "cristian.romero@sce.gob.ec"),
    "otac": ("Ec. Rosa Morales", "0998290196", "rosa.morales@sce.gob.ec"),
    "otal": ("Ec. Salomé Rosales", "0995041957", "salome.rosales@sce.gob.ec"),
}
ACTORES_ADICIONALES_GUAYAQUIL = ["José Matías"]

ARCHIVO_PRECARGA = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "congreso"
    / "LISTA DE INVITADOS 8sept.xlsx"
)
HASH_ARCHIVO_PRECARGA = (
    "ad7d6cf78428b46c505cd2bf31f77b47bc8f99ed20fb32c450d595f1a3fe3efa"
)
DISTRIBUCION_PRECARGA = {
    "guayaquil": 33,
    "manabi": 11,
    "cuenca": 10,
    "loja": 10,
    "sin_asignar": 7,
}

# Nuevos números y asignaciones incorporados en la versión del 13 de septiembre
# de 2026 de "LISTA DE INVITADOS 8sept con No oficios IR". Se conserva la fila
# de origen para completar los registros existentes sin duplicarlos.
ACTUALIZACIONES_OFICIOS_IR = [
    (3, "SCE-IGT-IR-2026-162", "Invitación general", "Milka"),
    (7, "SCE-IGT-IR-2026-163", "Invitación general", "OTAP"),
    (8, "SCE-IGT-IR-2026-164", "Invitación general", "Milka"),
    (9, "SCE-IGT-IR-2026-165", "Invitación general", "Milka"),
    (10, "SCE-IGT-IR-2026-166", "Invitación general", "OTAC"),
    (11, "SCE-IGT-IR-2026-167", "Invitación general", "Milka"),
    (12, "SCE-IGT-IR-2026-168", "Invitación general", "OTAC"),
    (13, "SCE-IGT-IR-2026-169", "Invitación general", "Milka"),
    (14, "SCE-IGT-IR-2026-170", "Invitación general", "Milka"),
    (15, "SCE-IGT-IR-2026-171", "Invitación general", "Milka"),
    (16, "SCE-IGT-IR-2026-172", "Invitación general", "Carlos G"),
    (38, "SCE-IGT-IR-2026-173", "Invitación general", "OTAL"),
    (40, "SCE-IGT-IR-2026-174", "Invitación a universidades", "Carlos G"),
    (41, "SCE-IGT-IR-2026-175", "Invitación a universidades", "Carlos G"),
    (62, "SCE-IGT-IR-2026-176", "Invitación general", "OTAC"),
    (65, "SCE-IGT-IR-2026-177", "Invitación general", "Carlos G"),
    (66, "SCE-IGT-IR-2026-178", "Invitación a universidades", "Milka"),
    (67, "SCE-IGT-IR-2026-179", "Invitación a universidades", "Milka"),
    (68, "SCE-IGT-IR-2026-180", "Invitación a universidades", "Milka"),
    (69, "SCE-IGT-IR-2026-181", "Invitación a universidades", "Carlos G"),
    (70, "SCE-IGT-IR-2026-182", "Invitación a universidades", "Carlos G"),
]

# La fila 71 recibió asignación de seguimiento, aunque conserva el oficio 643.
ACTUALIZACIONES_ASIGNACION = [
    (71, "SCE-2026-643", "Expositor", "Milka"),
]

_ETIQUETAS_CAMPOS = {
    "oficina": "Oficina",
    "responsable_id": "Responsable principal",
    "nombre_asistente_delegado": "Asistentes o delegados",
    "cargos_asistentes_delegados": "Cargos de asistentes o delegados",
    "confirmado": "Confirmado",
    "asistencia_21": "21 de octubre de 2026",
    "asistencia_22": "22 de octubre de 2026",
    "observaciones_seguimiento": "Observaciones de seguimiento",
    "numero_oficio": "N.º Oficio",
    "telefonos_institucionales": "Teléfono(s) institucional(es)",
    "tipo_invitacion": "Tipo de invitación",
    "observaciones_cruce": "Observaciones cruce de listado vs oficios generados",
    "proyeccion": "Proyección estudiantes",
    "confirmados": "Confirmados estudiantes",
    "contacto_nombre": "Nombre de contacto",
    "contacto_celular": "Celular de contacto",
    "observaciones": "Observaciones",
    "institucion": "Institución",
    "nombres": "Nombres",
    "celular": "Celular",
    "correo": "Correo",
    "activo": "Activo",
}

PROYECCION_ESTUDIANTES_INICIAL = [
    ("Autoridades", 15, 0, None),
    ("UG", 40, 0, "BUS"),
    ("Tecnológicos Gye", 70, 0, None),
    ("UNEMI", 50, 0, "BUS"),
    ("UTEG", 50, 0, None),
    ("UPSE", 50, 0, "BUS"),
    ("CEG", 30, 0, None),
    ("UCSG", 50, 0, None),
    ("TES", 40, 0, None),
    ("ESPOL", 40, 60, "BUS"),
    ("UTPL", 40, 0, "Salomé"),
    ("UPS", 40, 0, "María Morocho"),
    ("ECOTEC", 40, 30, None),
]


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _normalizar(valor: Any) -> str:
    texto = unicodedata.normalize("NFKD", _texto(valor).casefold())
    return " ".join("".join(c for c in texto if not unicodedata.combining(c)).split())


def precargar_proyeccion_estudiantes_desde_repositorio() -> int:
    """Crea las filas aprobadas sin reemplazar registros ya administrados."""
    filas = [
        {
            "clave_precarga": f"proyeccion_2026_{orden:02d}",
            "orden": orden,
            "institucion": institucion,
            "institucion_normalizada": _normalizar(institucion),
            "proyeccion": proyeccion,
            "confirmados": confirmados,
            "nota_original": nota,
        }
        for orden, (institucion, proyeccion, confirmados, nota) in enumerate(
            PROYECCION_ESTUDIANTES_INICIAL, start=1
        )
    ]
    with get_connection() as con:
        return precargar_proyeccion_estudiantes(con, filas)


def _entero_no_negativo(valor: Any, etiqueta: str) -> int:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        raise ValueError(f"{etiqueta} debe ser un número entero igual o mayor que cero.")
    try:
        numero = float(valor)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"{etiqueta} debe ser un número entero igual o mayor que cero."
        ) from exc
    if not numero.is_integer() or numero < 0:
        raise ValueError(f"{etiqueta} debe ser un número entero igual o mayor que cero.")
    return int(numero)


def _texto_editor(valor: Any) -> str:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    return str(valor).strip()


def _ultima_actualizacion_proyeccion(item: dict[str, Any]) -> str:
    if item.get("ultimo_actor_nombre"):
        fecha = _fecha_hora_ecuador(item.get("fecha_actualizacion"))
        detalle = str(item["ultimo_actor_nombre"])
        if fecha is not None:
            detalle += f" · {fecha:%d/%m/%Y %H:%M}"
        return detalle
    return ""


def _estado(valor: Any) -> str:
    normalizado = _normalizar(valor)
    if normalizado in {"si", "x", "confirmado", "confirmada"}:
        return "Sí"
    if normalizado in {"no", "rechazado", "rechazada"}:
        return "No"
    return "Pendiente"


def _indice_encabezado(encabezados: list[Any], nombre: str) -> int | None:
    buscado = _normalizar(nombre)
    for indice, encabezado in enumerate(encabezados, start=1):
        if _normalizar(encabezado) == buscado:
            return indice
    return None


_PATRON_TELEFONO = re.compile(
    r"(?:"
    r"\(?593(?:-\d)?\)?[\s-]*(?:\d{1,2}[\s-]*)?\d{3}[\s-]\d{4}"
    r"|\+593[\s-]*(?:\d{1,2}[\s-]*)?\d{3}[\s-]\d{4}"
    r"|\(0\d\)[\s-]*\d{3}[\s-]\d{4}"
    r"|0\d[\s-]+\d{3}[\s-]\d{4}"
    r"|1800[\s-]\d{3}[\s-]\d{3}"
    r"|(?<![\w@])\d{3}-\d{4}(?![\w.])"
    r"|\bext\.?\s*\d+\b"
    r")",
    re.IGNORECASE,
)


def _extraer_telefonos(valor: Any) -> str | None:
    """Extrae teléfonos de la celda de correo sin modificar el dato original."""
    texto = _texto(valor)
    if not texto:
        return None
    sin_correos = re.sub(r"[^\s,;|]+@[^\s,;|]+", " ", texto)
    encontrados: list[str] = []
    for coincidencia in _PATRON_TELEFONO.finditer(sin_correos):
        telefono = " ".join(coincidencia.group(0).split()).strip(" |,;")
        if telefono and telefono.casefold() not in {item.casefold() for item in encontrados}:
            encontrados.append(telefono)
    return "; ".join(encontrados) or None


def _separar_oficios_y_categorias(
    valor: Any,
) -> list[tuple[str | None, str | None]]:
    """Devuelve una invitación independiente por cada código de oficio."""
    original = _texto(valor)
    codigos_ir = re.findall(r"SCE-IGT-IR-2026-\d{3}", original, re.IGNORECASE)
    if codigos_ir:
        return [(codigo.upper(), None) for codigo in dict.fromkeys(codigos_ir)]
    numeros = re.findall(r"(?<!\d)(6(?:4[1-9]|5\d|6\d|7\d|8[0-6]))(?!\d)", original)
    if not numeros:
        return [(original or None, None)]
    numeros = list(dict.fromkeys(numeros))
    resultado: list[tuple[str, str]] = []
    for numero_texto in numeros:
        numero = int(numero_texto)
        categoria = (
            "Expositor" if 641 <= numero <= 644
            else "Invitación a universidades" if 645 <= numero <= 663
            else "Invitación general"
        )
        resultado.append((f"SCE-2026-{numero_texto}", categoria))
    return resultado


def analizar_excel_congreso(contenido: bytes) -> dict[str, Any]:
    """Valida y transforma el Excel sin escribir en la base de datos."""
    libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro["Invitados"] if "Invitados" in libro.sheetnames else libro.active
    libro_formulas = load_workbook(BytesIO(contenido), read_only=True, data_only=False)
    hoja_formulas = (
        libro_formulas["Invitados"]
        if "Invitados" in libro_formulas.sheetnames
        else libro_formulas.active
    )
    encabezados = [hoja.cell(2, col).value for col in range(1, hoja.max_column + 1)]

    nombres_requeridos = [
        "No", "Institucion", "Tipo institución", "Destinatario oficio", "Firma",
        "Nombre asistente o delegado", "Calidad", "Responsable seguimiento",
        "Confirmado", "21 de octubre de 2026", "22 de octubre de 2026", "Cargo",
        "Dirección", "Correo electrónico", "Sitio web", "Observaciones",
        "Responsable: Nombres", "Responsable: celular", "Responsable: Correo",
        "N° Oficio", "Observaciones cruce de listado vs oficios generados",
    ]
    columnas = {nombre: _indice_encabezado(encabezados, nombre) for nombre in nombres_requeridos}
    faltantes = [nombre for nombre, indice in columnas.items() if indice is None]
    if faltantes:
        return {
            "invitados": [], "responsables": [],
            "errores": ["Faltan columnas: " + ", ".join(faltantes)],
            "duplicados": [], "distribucion": {},
        }

    errores: list[str] = []
    duplicados: list[str] = []
    invitados: list[dict[str, Any]] = []
    responsables_por_clave: dict[tuple[str, str], dict[str, Any]] = {}
    claves_invitados: set[tuple[str, str, str]] = set()

    def valor(fila: int, columna: str) -> Any:
        return hoja.cell(fila, columnas[columna]).value

    for fila in range(3, hoja.max_row + 1):
        institucion = _texto(valor(fila, "Institucion"))
        numero_lista = _texto(valor(fila, "No"))
        tiene_datos_lista = any(
            _texto(valor(fila, campo))
            for campo in (
                "Tipo institución", "Destinatario oficio",
                "Responsable seguimiento", "N° Oficio",
            )
        )
        if not institucion or not tiene_datos_lista:
            continue
        # Algunos editores no guardan el resultado calculado de la fórmula de
        # numeración. Solo en ese caso la posición conserva el valor estable;
        # filas auxiliares realmente vacías continúan excluidas.
        if not numero_lista:
            formula_numero = hoja_formulas.cell(fila, columnas["No"]).value
            if not (isinstance(formula_numero, str) and formula_numero.startswith("=")):
                continue
            numero_lista = str(fila - 2)

        etiqueta_responsable = _normalizar(valor(fila, "Responsable seguimiento"))
        oficina = MAPEO_RESPONSABLE_OFICINA.get(etiqueta_responsable)
        if etiqueta_responsable and oficina is None:
            errores.append(
                f"Fila {fila}: responsable de seguimiento no reconocido: "
                f"{_texto(valor(fila, 'Responsable seguimiento'))}."
            )

        nombre_responsable = _texto(valor(fila, "Responsable: Nombres"))
        celular_responsable = _texto(valor(fila, "Responsable: celular"))
        correo_responsable = _texto(valor(fila, "Responsable: Correo")).lower()
        datos_responsable = MAPEO_RESPONSABLE_DATOS.get(etiqueta_responsable)
        if oficina and datos_responsable:
            nombre_responsable = nombre_responsable or datos_responsable[0]
            celular_responsable = celular_responsable or datos_responsable[1]
            correo_responsable = correo_responsable or datos_responsable[2]
        if oficina and nombre_responsable:
            clave_responsable = (oficina, nombre_responsable.casefold())
            responsable = {
                "oficina": oficina, "nombres": nombre_responsable,
                "celular": celular_responsable, "correo": correo_responsable,
            }
            anterior = responsables_por_clave.get(clave_responsable)
            if anterior and anterior != responsable:
                errores.append(
                    f"Fila {fila}: {nombre_responsable} aparece con datos de contacto diferentes."
                )
            else:
                responsables_por_clave[clave_responsable] = responsable
            if not _celular_valido(celular_responsable):
                errores.append(
                    f"Fila {fila}: {nombre_responsable} no tiene un celular válido."
                )
            if not _correo_valido(correo_responsable):
                errores.append(
                    f"Fila {fila}: {nombre_responsable} no tiene un correo válido."
                )
        elif oficina and not nombre_responsable:
            errores.append(f"Fila {fila}: la asignación no tiene nombre de responsable.")

        destinatario = _texto(valor(fila, "Destinatario oficio"))
        numero_oficio_original = _texto(valor(fila, "N° Oficio"))
        oficios = _separar_oficios_y_categorias(
            numero_oficio_original
        )
        calidad = _texto(valor(fila, "Calidad"))
        tipo_desde_calidad = {
            "invitacion general": "Invitación general",
            "invitacion universidad": "Invitación a universidades",
            "expositor": "Expositor",
        }.get(_normalizar(calidad))
        invitado_base = {
                "fila_origen": fila,
                "numero_lista": numero_lista,
                "institucion": institucion,
                "tipo_institucion": _texto(valor(fila, "Tipo institución")) or None,
                "destinatario_oficio": destinatario or None,
                "firma": _texto(valor(fila, "Firma")) or None,
                "calidad": calidad or None,
                "cargo": _texto(valor(fila, "Cargo")) or None,
                "direccion": _texto(valor(fila, "Dirección")) or None,
                "correo_institucional": _texto(valor(fila, "Correo electrónico")) or None,
                "telefonos_institucionales": _extraer_telefonos(
                    valor(fila, "Correo electrónico")
                ),
                "sitio_web": _texto(valor(fila, "Sitio web")) or None,
                "oficina": oficina,
                "_responsable_nombre": nombre_responsable or None,
                "nombre_asistente_delegado": _texto(
                    valor(fila, "Nombre asistente o delegado")
                ) or None,
                "confirmado": _estado(valor(fila, "Confirmado")),
                "asistencia_21": _estado(valor(fila, "21 de octubre de 2026")),
                "asistencia_22": _estado(valor(fila, "22 de octubre de 2026")),
                "observaciones_seguimiento": _texto(valor(fila, "Observaciones")) or None,
                "observaciones_cruce": _texto(
                    valor(fila, "Observaciones cruce de listado vs oficios generados")
                ) or None,
            }
        for numero_oficio, tipo_invitacion in oficios:
            tipo_invitacion = tipo_invitacion or tipo_desde_calidad
            clave_invitado = (
                _normalizar(institucion),
                _normalizar(destinatario),
                _normalizar(numero_oficio),
            )
            if clave_invitado in claves_invitados:
                duplicados.append(
                    f"Fila {fila}: {institucion} / {destinatario or 'sin destinatario'}"
                )
            claves_invitados.add(clave_invitado)
            invitado = dict(invitado_base)
            invitado["numero_oficio"] = numero_oficio
            invitado["tipo_invitacion"] = tipo_invitacion
            invitados.append(invitado)

    distribucion = Counter(item["oficina"] or "sin_asignar" for item in invitados)
    return {
        "invitados": invitados,
        "responsables": list(responsables_por_clave.values()),
        "errores": list(dict.fromkeys(errores)),
        "duplicados": list(dict.fromkeys(duplicados)),
        "distribucion": dict(distribucion),
    }


def precargar_congreso_desde_repositorio() -> int:
    """Importa el archivo aprobado una sola vez cuando la base está vacía."""
    with get_connection() as con:
        if contar_invitados_congreso(con) > 0:
            return 0

    contenido = ARCHIVO_PRECARGA.read_bytes()
    hash_archivo = sha256(contenido).hexdigest()
    if hash_archivo != HASH_ARCHIVO_PRECARGA:
        raise ValueError("El archivo de precarga no coincide con la versión aprobada.")

    resultado = analizar_excel_congreso(contenido)
    if resultado["errores"] or resultado["duplicados"]:
        detalles = [*resultado["errores"], *resultado["duplicados"]]
        raise ValueError("La precarga contiene incidencias: " + " | ".join(detalles))
    if len(resultado["invitados"]) != 71:
        raise ValueError("La precarga debe contener exactamente 71 invitaciones.")
    if len(resultado["responsables"]) != 5:
        raise ValueError("La precarga debe contener exactamente 5 responsables.")
    if resultado["distribucion"] != DISTRIBUCION_PRECARGA:
        raise ValueError(
            "La distribución de oficinas del archivo de precarga no es la esperada."
        )

    try:
        with get_connection() as con:
            return importar_datos_congreso(
                con,
                [dict(item) for item in resultado["invitados"]],
                resultado["responsables"],
                ARCHIVO_PRECARGA.name,
                hash_archivo,
                "Carga inicial — Guayaquil",
            )
    except ValueError:
        # Otra instancia puede haber terminado la carga mientras esta esperaba el bloqueo.
        with get_connection() as con:
            if contar_invitados_congreso(con) > 0:
                return 0
        raise


NUEVOS_OFICIOS_FIRMADOS = [
    {
        "numero_lista": None,
        "institucion": "UTEQ — Universidad Técnica Estatal de Quevedo",
        "tipo_institucion": "Universidad",
        "destinatario_oficio": "Yenny Torres Navarrete, PhD.",
        "firma": "Intendente Regional de Guayaquil",
        "calidad": "Invitación Universidad",
        "cargo": "Rectora",
        "direccion": (
            "Campus Central, Av. Quito Km. 1½ vía a Santo Domingo de los "
            "Tsáchilas, Quevedo"
        ),
        "correo_institucional": "info@uteq.edu.ec",
        "telefonos_institucionales": None,
        "sitio_web": None,
        "tipo_invitacion": "Invitación a universidades",
        "oficina": "guayaquil",
        "numero_oficio": "SCE-IGT-IR-2026-154",
    },
    {
        "numero_lista": None,
        "institucion": (
            "UNIANDES — Universidad Regional Autónoma de los Andes, Extensión Quevedo"
        ),
        "tipo_institucion": "Universidad",
        "destinatario_oficio": "Danilo Viteri Intriago, PhD.",
        "firma": "Intendente Regional de Guayaquil",
        "calidad": "Invitación Universidad",
        "cargo": "Director",
        "direccion": (
            "Vía a Valencia Km. 5½, Campus Universitario Dr. Gustavo Álvarez "
            "Gavilanes, Quevedo"
        ),
        "correo_institucional": (
            "direccionquevedo@uniandes.edu.ec; uq.rb.secretariader@uniandes.edu.ec; "
            "uq.derecho@uniandes.edu.ec; uq.asissecretaria@uniandes.edu.ec"
        ),
        "telefonos_institucionales": None,
        "sitio_web": None,
        "tipo_invitacion": "Invitación a universidades",
        "oficina": "guayaquil",
        "numero_oficio": "SCE-IGT-IR-2026-155",
    },
]


def sincronizar_congreso_desde_documentos() -> int:
    """Aplica una sola vez teléfonos, códigos completos, categorías y nuevos oficios."""
    resultado = analizar_excel_congreso(ARCHIVO_PRECARGA.read_bytes())
    actualizaciones = [
        {
            "fila_origen": item["fila_origen"],
            "telefonos_institucionales": item.get("telefonos_institucionales"),
            "tipo_invitacion": item.get("tipo_invitacion"),
            "numero_oficio": item.get("numero_oficio"),
        }
        for item in resultado["invitados"]
    ]
    por_fila = {item["fila_origen"]: item for item in actualizaciones}
    for fila, numero_oficio, tipo_invitacion, etiqueta_responsable in [
        *ACTUALIZACIONES_OFICIOS_IR,
        *ACTUALIZACIONES_ASIGNACION,
    ]:
        item = por_fila[fila]
        item.update(
            numero_oficio=numero_oficio,
            tipo_invitacion=tipo_invitacion,
            oficina=MAPEO_RESPONSABLE_OFICINA[etiqueta_responsable.casefold()],
            responsable_nombre=MAPEO_RESPONSABLE_DATOS[
                etiqueta_responsable.casefold()
            ][0],
        )
    with get_connection() as con:
        return sincronizar_documentos_congreso(
            con, actualizaciones, NUEVOS_OFICIOS_FIRMADOS
        )


def _correo_valido(correo: str) -> bool:
    return bool(re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", correo.strip()))


def _celular_valido(celular: str) -> bool:
    return len(re.sub(r"\D", "", celular)) >= 7


def _nombre_oficina(oficina: str | None) -> str:
    return OFICINAS.get(oficina, "Sin asignar")


def _firmado_por(valor: Any) -> str:
    firma = _normalizar(valor)
    if firma == "ir" or "intendente regional" in firma:
        return "IR"
    if "superintendente" in firma:
        return "Superintendente"
    return ""


def _fecha_hora_ecuador(valor: Any) -> datetime | None:
    if valor is None or valor == "":
        return None
    fecha = valor.to_pydatetime() if isinstance(valor, pd.Timestamp) else valor
    if not isinstance(fecha, datetime):
        fecha = pd.to_datetime(fecha).to_pydatetime()
    if fecha.tzinfo is None:
        fecha = fecha.replace(tzinfo=timezone.utc)
    return fecha.astimezone(ZONA_HORARIA_ECUADOR).replace(tzinfo=None)


def _ultimo_cambio_invitado(item: dict[str, Any]) -> str:
    fecha = _fecha_hora_ecuador(item.get("ultimo_cambio_fecha"))
    if fecha is None:
        return "Sin cambios"
    actor = item.get("ultimo_cambio_actor") or "Sistema"
    campo = _ETIQUETAS_CAMPOS.get(
        item.get("ultimo_cambio_campo"),
        item.get("ultimo_cambio_campo") or item.get("ultimo_cambio_accion") or "Cambio",
    )
    valor = _texto(item.get("ultimo_cambio_valor"))
    detalle = f"{campo}: {valor}" if valor else campo
    return f"{fecha:%d/%m/%Y %H:%M} · {actor} · {detalle}"


def _tabla_asistentes(nombres: Any, cargos: Any) -> pd.DataFrame:
    lista_nombres = [linea.strip() for linea in _texto(nombres).splitlines() if linea.strip()]
    lista_cargos = [linea.strip() for linea in _texto(cargos).splitlines()]
    cantidad = max(len(lista_nombres), len(lista_cargos), 1)
    lista_nombres.extend([""] * (cantidad - len(lista_nombres)))
    lista_cargos.extend([""] * (cantidad - len(lista_cargos)))
    return pd.DataFrame({"Nombre": lista_nombres, "Cargo": lista_cargos})


def _serializar_asistentes(tabla: pd.DataFrame) -> tuple[str, str]:
    nombres: list[str] = []
    cargos: list[str] = []
    for _, fila in tabla.fillna("").iterrows():
        nombre = _texto(fila.get("Nombre"))
        cargo = _texto(fila.get("Cargo"))
        if not nombre and cargo:
            raise ValueError("Cada cargo debe estar asociado a un nombre.")
        if nombre:
            nombres.append(nombre)
            cargos.append(cargo)
    return "\n".join(nombres), "\n".join(cargos)


def _datos_exportacion(
    invitados: list[dict[str, Any]], historial: list[dict[str, Any]]
) -> bytes:
    columnas_invitados = {
        "numero_lista": "No", "institucion": "Institución",
        "tipo_institucion": "Tipo institución", "destinatario_oficio": "Destinatario oficio",
        "firma": "Firmado por", "nombre_asistente_delegado": "Asistentes o delegados",
        "cargos_asistentes_delegados": "Cargos de asistentes o delegados",
        "calidad": "Calidad", "oficina": "Oficina",
        "responsable_nombres": "Responsable: Nombres",
        "responsable_celular": "Responsable: celular",
        "responsable_correo": "Responsable: Correo", "confirmado": "Confirmado",
        "asistencia_21": "21 de octubre de 2026",
        "asistencia_22": "22 de octubre de 2026", "cargo": "Cargo",
        "direccion": "Dirección", "correo_institucional": "Correo electrónico",
        "telefonos_institucionales": "Teléfono(s) institucional(es)",
        "sitio_web": "Sitio web", "observaciones_seguimiento": "Observaciones",
        "tipo_invitacion": "Tipo de invitación",
        "numero_oficio": "N.º Oficio",
        "observaciones_cruce": "Observaciones cruce de listado vs oficios generados",
        "fecha_actualizacion": "Última actualización (Ecuador)",
    }
    filas_invitados = []
    for item in invitados:
        fila = {etiqueta: item.get(campo) for campo, etiqueta in columnas_invitados.items()}
        fila["Oficina"] = _nombre_oficina(item.get("oficina"))
        fila["Firmado por"] = _firmado_por(item.get("firma"))
        fila["Última actualización (Ecuador)"] = _fecha_hora_ecuador(
            item.get("fecha_actualizacion")
        )
        filas_invitados.append(fila)

    columnas_historial = {
        "fecha_cambio": "Fecha (Ecuador)", "institucion": "Institución", "oficina": "Oficina",
        "actor_nombre": "Actualizado por", "accion": "Acción", "campo": "Campo",
        "valor_anterior": "Valor anterior", "valor_nuevo": "Valor nuevo",
    }
    filas_historial = []
    for item in historial:
        fila = {etiqueta: item.get(campo) for campo, etiqueta in columnas_historial.items()}
        fila["Fecha (Ecuador)"] = _fecha_hora_ecuador(item.get("fecha_cambio"))
        fila["Oficina"] = _nombre_oficina(item.get("oficina"))
        fila["Campo"] = _ETIQUETAS_CAMPOS.get(item.get("campo"), item.get("campo"))
        filas_historial.append(fila)

    salida = BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        pd.DataFrame(filas_invitados, columns=columnas_invitados.values()).to_excel(
            writer, sheet_name="Invitados", index=False
        )
        pd.DataFrame(filas_historial, columns=columnas_historial.values()).to_excel(
            writer, sheet_name="Historial", index=False
        )
        for hoja in writer.book.worksheets:
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
            for celda in hoja[1]:
                celda.fill = PatternFill("solid", fgColor=COLOR_AZUL.removeprefix("#"))
                celda.font = Font(color="FFFFFF", bold=True)
                celda.alignment = Alignment(horizontal="center", vertical="center")
            for columna in hoja.columns:
                ancho = min(max(len(_texto(celda.value)) for celda in columna) + 2, 48)
                hoja.column_dimensions[columna[0].column_letter].width = max(ancho, 12)
    return salida.getvalue()


def _actor_actual(oficina_id: str) -> tuple[int | None, str]:
    with get_connection() as con:
        responsables = [
            dict(item)
            for item in listar_responsables_congreso(con, oficina_id, solo_activos=True)
        ]
    opciones: list[tuple[str, int | str]] = [
        ("responsable", item["id"]) for item in responsables
    ]
    nombres_opciones = {
        ("responsable", item["id"]): item["nombres"] for item in responsables
    }
    if oficina_id == "guayaquil":
        for nombre in ACTORES_ADICIONALES_GUAYAQUIL:
            opcion = ("adicional", nombre)
            opciones.append(opcion)
            nombres_opciones[opcion] = nombre
    if not opciones:
        st.info("Agrega un responsable en la pestaña Responsables para registrar cambios.")
        return None, ""
    opcion_actor = st.selectbox(
        "¿Quién está actualizando la información?", opciones,
        format_func=lambda valor: nombres_opciones[valor],
        key=f"congreso_actor_{oficina_id}",
    )
    actor_id = int(opcion_actor[1]) if opcion_actor[0] == "responsable" else None
    return actor_id, nombres_opciones[opcion_actor]


def _mostrar_indicadores(invitados: list[dict[str, Any]]) -> None:
    valores = (
        len(invitados),
        sum(item["confirmado"] == "Sí" for item in invitados),
        sum(item["confirmado"] == "No" for item in invitados),
        sum(item["confirmado"] == "Pendiente" for item in invitados),
        sum(item.get("responsable_id") is None for item in invitados),
        sum(item["asistencia_21"] == "Sí" for item in invitados),
        sum(item["asistencia_22"] == "Sí" for item in invitados),
    )
    cols = st.columns(7)
    etiquetas = (
        "Invitados", "Confirmados", "No asistirán", "Pendientes",
        "Sin responsable", "Asisten 21", "Asisten 22",
    )
    for col, etiqueta, valor in zip(cols, etiquetas, valores):
        col.metric(etiqueta, valor)


def _mostrar_avance_oficinas(invitados: list[dict[str, Any]]) -> None:
    filas = []
    for oficina in ["guayaquil", "manabi", "cuenca", "loja", None]:
        grupo = [item for item in invitados if item.get("oficina") == oficina]
        if not grupo:
            continue
        filas.append(
            {
                "Oficina": _nombre_oficina(oficina),
                "Invitados": len(grupo),
                "Confirmados": sum(item["confirmado"] == "Sí" for item in grupo),
                "No asistirán": sum(item["confirmado"] == "No" for item in grupo),
                "Pendientes": sum(item["confirmado"] == "Pendiente" for item in grupo),
                "Sin responsable": sum(item.get("responsable_id") is None for item in grupo),
                "Asisten 21": sum(item["asistencia_21"] == "Sí" for item in grupo),
                "Asisten 22": sum(item["asistencia_22"] == "Sí" for item in grupo),
            }
        )
    with st.expander("Avance por oficina", expanded=True):
        st.dataframe(pd.DataFrame(filas), hide_index=True, use_container_width=True)


def _vista_seguimiento(
    oficina_id: str, es_master: bool, actor_id: int | None, actor_nombre: str
) -> None:
    with get_connection() as con:
        datos = [dict(item) for item in listar_invitados_congreso(
            con, None if es_master else oficina_id
        )]
    if not datos:
        st.info("Todavía no existen invitados. Realiza la carga inicial desde Guayaquil.")
        return

    _mostrar_indicadores(datos)
    if es_master:
        _mostrar_avance_oficinas(datos)
    st.divider()
    c1, c2, c3 = st.columns([1.3, 1, 2])
    with c1:
        filtro_oficina = "Todas"
        if es_master:
            filtro_oficina = st.selectbox(
                "Oficina", ["Todas", "Sin asignar", *OFICINAS.values()],
                key="congreso_filtro_oficina",
            )
    with c2:
        filtro_estado = st.selectbox(
            "Confirmación", ["Todos", *ESTADOS], key="congreso_filtro_estado"
        )
    with c3:
        busqueda = st.text_input(
            "Buscar institución o destinatario", key="congreso_busqueda"
        ).strip().casefold()

    filtrados = datos
    if es_master and filtro_oficina != "Todas":
        oficina_filtro = None if filtro_oficina == "Sin asignar" else next(
            clave for clave, nombre in OFICINAS.items() if nombre == filtro_oficina
        )
        filtrados = [item for item in filtrados if item.get("oficina") == oficina_filtro]
    if filtro_estado != "Todos":
        filtrados = [item for item in filtrados if item["confirmado"] == filtro_estado]
    if busqueda:
        filtrados = [
            item for item in filtrados
            if busqueda in _texto(item.get("institucion")).casefold()
            or busqueda in _texto(item.get("destinatario_oficio")).casefold()
        ]

    # La lista maestra se conserva intacta. En la tabla de trabajo solo se
    # muestran casos que ya cuentan con número de oficio.
    filtrados = [item for item in filtrados if _texto(item.get("numero_oficio"))]

    with st.expander("Filtros por columna", expanded=False):
        st.caption("Combina los filtros para acotar los casos que deseas revisar.")
        f1, f2, f3, f4 = st.columns(4)
        filtro_institucion = f1.text_input("Institución", key="congreso_col_institucion")
        filtro_responsable = f2.multiselect(
            "Responsable",
            sorted({_texto(item.get("responsable_nombres") or "Sin asignar") for item in filtrados}),
            key="congreso_col_responsable",
        )
        filtro_confirmado = f3.multiselect(
            "Confirmado", ESTADOS, key="congreso_col_confirmado"
        )
        filtro_oficio = f4.text_input("N.º Oficio", key="congreso_col_oficio")
        f5, f6, f7, f8 = st.columns(4)
        filtro_oficina_col = f5.multiselect(
            "Oficina",
            sorted({_nombre_oficina(item.get("oficina")) for item in filtrados}),
            key="congreso_col_oficina",
        )
        filtro_21 = f6.multiselect("21 oct.", ESTADOS, key="congreso_col_21")
        filtro_22 = f7.multiselect("22 oct.", ESTADOS, key="congreso_col_22")
        filtro_delegado = f8.text_input("Asistentes o delegados", key="congreso_col_delegado")
        f9, f10, f11, f12 = st.columns(4)
        filtro_tipo = f9.multiselect(
            "Tipo de invitación",
            sorted({_texto(item.get("tipo_invitacion")) for item in filtrados if _texto(item.get("tipo_invitacion"))}),
            key="congreso_col_tipo_invitacion",
        )
        filtro_cargo = f10.text_input("Cargos", key="congreso_col_cargos")
        filtro_firma = f11.multiselect(
            "Firmado por",
            sorted({_firmado_por(item.get("firma")) for item in filtrados if _firmado_por(item.get("firma"))}),
            key="congreso_col_firmado_por",
        )
        filtro_historial = f12.text_input(
            "Historial/último cambio", key="congreso_col_historial"
        )

    if filtro_institucion.strip():
        valor = filtro_institucion.strip().casefold()
        filtrados = [item for item in filtrados if valor in _texto(item.get("institucion")).casefold()]
    if filtro_responsable:
        filtrados = [item for item in filtrados if _texto(item.get("responsable_nombres") or "Sin asignar") in filtro_responsable]
    if filtro_confirmado:
        filtrados = [item for item in filtrados if item.get("confirmado") in filtro_confirmado]
    if filtro_oficio.strip():
        valor = filtro_oficio.strip().casefold()
        filtrados = [item for item in filtrados if valor in _texto(item.get("numero_oficio")).casefold()]
    if filtro_oficina_col:
        filtrados = [item for item in filtrados if _nombre_oficina(item.get("oficina")) in filtro_oficina_col]
    if filtro_21:
        filtrados = [item for item in filtrados if item.get("asistencia_21") in filtro_21]
    if filtro_22:
        filtrados = [item for item in filtrados if item.get("asistencia_22") in filtro_22]
    if filtro_delegado.strip():
        valor = filtro_delegado.strip().casefold()
        filtrados = [item for item in filtrados if valor in _texto(item.get("nombre_asistente_delegado")).casefold()]
    if filtro_tipo:
        filtrados = [item for item in filtrados if _texto(item.get("tipo_invitacion")) in filtro_tipo]
    if filtro_cargo.strip():
        valor = filtro_cargo.strip().casefold()
        filtrados = [item for item in filtrados if valor in _texto(item.get("cargos_asistentes_delegados")).casefold()]
    if filtro_firma:
        filtrados = [item for item in filtrados if _firmado_por(item.get("firma")) in filtro_firma]
    if filtro_historial.strip():
        valor = filtro_historial.strip().casefold()
        filtrados = [
            item for item in filtrados
            if valor in _ultimo_cambio_invitado(item).casefold()
        ]

    tabla = pd.DataFrame([
        {
            "Institución": item["institucion"], "Oficina": _nombre_oficina(item.get("oficina")),
            "Responsable": item.get("responsable_nombres") or "Sin asignar",
            "Tipo de invitación": item.get("tipo_invitacion") or "",
            "Confirmado": item["confirmado"], "21 oct.": item["asistencia_21"],
            "22 oct.": item["asistencia_22"],
            "Asistentes o delegados": (
                item.get("nombre_asistente_delegado") or ""
            ).replace("\n", "; "),
            "Cargos": (item.get("cargos_asistentes_delegados") or "").replace("\n", "; "),
            "Firmado por": _firmado_por(item.get("firma")),
            "N.º Oficio": item.get("numero_oficio") or "",
            "Historial/último cambio": _ultimo_cambio_invitado(item),
        }
        for item in filtrados
    ])
    st.dataframe(
        tabla,
        hide_index=True,
        use_container_width=True,
        height=330,
        column_config={
            "Historial/último cambio": st.column_config.TextColumn(width="large")
        },
    )
    st.caption(f"{len(filtrados)} casos visibles. Los casos sin número de oficio permanecen guardados, pero no aparecen en esta tabla.")
    if not filtrados:
        st.info("No hay invitados que coincidan con los filtros.")
        return

    por_id = {item["id"]: item for item in filtrados}
    seleccionado_id = st.selectbox(
        "Selecciona un invitado para actualizar", list(por_id),
        format_func=lambda valor: (
            f"{por_id[valor]['institucion']} — "
            f"{por_id[valor].get('numero_oficio') or 'sin oficio'} — "
            f"{por_id[valor].get('destinatario_oficio') or 'sin destinatario'}"
        ),
        key="congreso_invitado_editar",
    )
    invitado = por_id[seleccionado_id]

    with st.expander("Datos originales del listado"):
        st.markdown(
            f"**Institución:** {invitado['institucion']}  \n"
            f"**Destinatario:** {invitado.get('destinatario_oficio') or '—'}  \n"
            f"**Cargo:** {invitado.get('cargo') or '—'}  \n"
            f"**Correo institucional:** {invitado.get('correo_institucional') or '—'}  \n"
            f"**Teléfono(s):** {invitado.get('telefonos_institucionales') or '—'}  \n"
            f"**Tipo de invitación:** {invitado.get('tipo_invitacion') or '—'}  \n"
            f"**Firmado por:** {_firmado_por(invitado.get('firma')) or '—'}  \n"
            f"**N.º Oficio:** {invitado.get('numero_oficio') or '—'}  \n"
            f"**Dirección:** {invitado.get('direccion') or '—'}"
        )

    oficina_destino = invitado.get("oficina")
    if es_master:
        opciones_oficina = [None, *OFICINAS]
        oficina_destino = st.selectbox(
            "Oficina responsable", opciones_oficina,
            index=opciones_oficina.index(oficina_destino),
            format_func=_nombre_oficina,
            key=f"congreso_oficina_{seleccionado_id}",
        )

    with get_connection() as con:
        responsables = [dict(item) for item in listar_responsables_congreso(
            con, oficina_destino, solo_activos=True
        )] if oficina_destino else []
    responsables_por_id = {item["id"]: item for item in responsables}
    opciones_responsable = [None, *responsables_por_id]
    actual_responsable = invitado.get("responsable_id")
    indice_responsable = opciones_responsable.index(actual_responsable) if actual_responsable in opciones_responsable else 0

    with st.form(f"form_congreso_invitado_{seleccionado_id}"):
        responsable_id = st.selectbox(
            "Responsable principal", opciones_responsable, index=indice_responsable,
            format_func=lambda valor: "Sin asignar" if valor is None else responsables_por_id[valor]["nombres"],
        )
        if responsable_id:
            responsable = responsables_por_id[responsable_id]
            st.caption(f"{responsable['celular']} · {responsable['correo']}")
        st.markdown("**Asistentes o delegados**")
        st.caption("Agrega una fila por persona e indica su cargo. Puedes añadir o eliminar filas.")
        tabla_asistentes = st.data_editor(
            _tabla_asistentes(
                invitado.get("nombre_asistente_delegado"),
                invitado.get("cargos_asistentes_delegados"),
            ),
            hide_index=True,
            num_rows="dynamic",
            use_container_width=True,
            column_config={
                "Nombre": st.column_config.TextColumn("Nombre asistente o delegado"),
                "Cargo": st.column_config.TextColumn("Cargo"),
            },
            key=f"congreso_asistentes_{seleccionado_id}",
        )
        c1, c2, c3 = st.columns(3)
        confirmado = c1.selectbox("Confirmado", ESTADOS, index=ESTADOS.index(invitado["confirmado"]))
        asistencia_21 = c2.selectbox(
            "21 de octubre de 2026", ESTADOS, index=ESTADOS.index(invitado["asistencia_21"])
        )
        asistencia_22 = c3.selectbox(
            "22 de octubre de 2026", ESTADOS, index=ESTADOS.index(invitado["asistencia_22"])
        )
        numero_oficio = st.text_input("N.º Oficio", value=invitado.get("numero_oficio") or "")
        observaciones = st.text_area(
            "Observaciones de seguimiento", value=invitado.get("observaciones_seguimiento") or ""
        )
        observaciones_cruce = st.text_area(
            "Observaciones cruce de listado vs oficios generados",
            value=invitado.get("observaciones_cruce") or "",
        )
        guardar = st.form_submit_button("Guardar cambios", type="primary", disabled=not actor_nombre)

    if guardar:
        try:
            asistentes, cargos_asistentes = _serializar_asistentes(tabla_asistentes)
            with get_connection() as con:
                cantidad = actualizar_invitado_congreso(
                    con, seleccionado_id,
                    {
                        "oficina": oficina_destino, "responsable_id": responsable_id,
                        "nombre_asistente_delegado": asistentes,
                        "cargos_asistentes_delegados": cargos_asistentes,
                        "confirmado": confirmado,
                        "asistencia_21": asistencia_21, "asistencia_22": asistencia_22,
                        "observaciones_seguimiento": observaciones,
                        "numero_oficio": numero_oficio,
                        "observaciones_cruce": observaciones_cruce,
                    },
                    actor_id, actor_nombre, oficina_id,
                    None if es_master else oficina_id,
                )
            if cantidad:
                st.success(f"Se guardaron {cantidad} cambios en el historial.")
                st.rerun()
            else:
                st.info("No se detectaron cambios.")
        except (ValueError, PermissionError) as exc:
            st.error(str(exc))
        except Exception as exc:
            st.error(f"No se pudieron guardar los cambios: {exc}")


def _vista_responsables(
    oficina_id: str, es_master: bool, actor_id: int | None, actor_nombre: str
) -> None:
    oficina_gestion = oficina_id
    if es_master:
        oficina_gestion = st.selectbox(
            "Oficina del directorio", list(OFICINAS), format_func=_nombre_oficina,
            key="congreso_directorio_oficina",
        )
    with get_connection() as con:
        responsables = [dict(item) for item in listar_responsables_congreso(con, oficina_gestion)]

    tabla = pd.DataFrame([
        {
            "Nombres": item["nombres"], "Celular": item["celular"],
            "Correo": item["correo"], "Activo": "Sí" if item["activo"] else "No",
            "Invitados asignados": item["invitados_asignados"],
        }
        for item in responsables
    ])
    st.dataframe(tabla, hide_index=True, use_container_width=True)

    with st.expander("Agregar coordinador o analista", expanded=not responsables):
        with st.form(f"form_nuevo_responsable_{oficina_gestion}", clear_on_submit=True):
            nombres = st.text_input("Nombres completos")
            celular = st.text_input("Celular")
            correo = st.text_input("Correo")
            agregar = st.form_submit_button("Agregar responsable", type="primary")
        if agregar:
            if not nombres.strip() or not _celular_valido(celular) or not _correo_valido(correo):
                st.error("Ingresa nombres, un celular válido y un correo válido.")
            else:
                try:
                    with get_connection() as con:
                        crear_responsable_congreso(
                            con, oficina_gestion, nombres, celular, correo, actor_id,
                            actor_nombre or f"Perfil {_nombre_oficina(oficina_id)}", oficina_id,
                        )
                    st.success("Responsable agregado.")
                    st.rerun()
                except Exception as exc:
                    st.error(f"No se pudo agregar el responsable: {exc}")

    if not responsables:
        return
    por_id = {item["id"]: item for item in responsables}
    editar_id = st.selectbox(
        "Responsable que deseas modificar", list(por_id),
        format_func=lambda valor: por_id[valor]["nombres"],
        key=f"congreso_responsable_editar_{oficina_gestion}",
    )
    actual = por_id[editar_id]
    with st.form(f"form_editar_responsable_{editar_id}"):
        nombres_editar = st.text_input("Nombres completos", value=actual["nombres"])
        celular_editar = st.text_input("Celular", value=actual["celular"])
        correo_editar = st.text_input("Correo", value=actual["correo"])
        activo = st.checkbox("Responsable activo", value=actual["activo"])
        guardar = st.form_submit_button("Guardar responsable", type="primary", disabled=not actor_nombre)
    if guardar:
        if not nombres_editar.strip() or not _celular_valido(celular_editar) or not _correo_valido(correo_editar):
            st.error("Ingresa nombres, un celular válido y un correo válido.")
            return
        try:
            with get_connection() as con:
                actualizar_responsable_congreso(
                    con, editar_id, nombres_editar, celular_editar, correo_editar, activo,
                    actor_id, actor_nombre, oficina_id, None if es_master else oficina_id,
                )
            st.success("Responsable actualizado.")
            st.rerun()
        except (ValueError, PermissionError) as exc:
            st.error(str(exc))
        except Exception as exc:
            st.error(f"No se pudo actualizar el responsable: {exc}")


def _vista_proyeccion_estudiantes(
    actor_id: int | None, actor_nombre: str
) -> None:
    with get_connection() as con:
        activos = [dict(item) for item in listar_proyeccion_estudiantes(con, True)]
        todos = [dict(item) for item in listar_proyeccion_estudiantes(con, None)]

    col1, col2, col3 = st.columns(3)
    col1.metric("Instituciones activas", len(activos))
    col2.metric("Estudiantes proyectados", sum(item["proyeccion"] for item in activos))
    col3.metric("Estudiantes confirmados", sum(item["confirmados"] for item in activos))
    st.caption(
        "Edita confirmados, datos de contacto y observaciones. Para cambiar "
        "confirmados, completa primero el nombre y celular del contacto."
    )

    if activos:
        tabla = pd.DataFrame(
            [
                {
                    "id": item["id"],
                    "Institución": item["institucion"],
                    "Proyección": item["proyeccion"],
                    "Confirmados": item["confirmados"],
                    "Nombre de contacto": item.get("contacto_nombre") or "",
                    "Celular de contacto": item.get("contacto_celular") or "",
                    "Observaciones": item.get("observaciones") or "",
                    "Última actualización": _ultima_actualizacion_proyeccion(item),
                }
                for item in activos
            ]
        ).set_index("id")
        with st.form("form_editar_proyeccion_estudiantes"):
            editada = st.data_editor(
                tabla,
                hide_index=True,
                use_container_width=True,
                height=min(560, 38 * (len(tabla) + 1)),
                disabled=["Institución", "Proyección", "Última actualización"],
                column_config={
                    "Institución": st.column_config.TextColumn("Institución", width="large"),
                    "Proyección": st.column_config.NumberColumn(
                        "Proyección", min_value=0, step=1, format="%d"
                    ),
                    "Confirmados": st.column_config.NumberColumn(
                        "Confirmados", min_value=0, step=1, format="%d"
                    ),
                    "Nombre de contacto": st.column_config.TextColumn(
                        "Nombre de contacto", width="medium"
                    ),
                    "Celular de contacto": st.column_config.TextColumn(
                        "Celular de contacto", width="medium"
                    ),
                    "Observaciones": st.column_config.TextColumn(
                        "Observaciones", width="large"
                    ),
                    "Última actualización": st.column_config.TextColumn(
                        "Última actualización", width="large"
                    ),
                },
                key="congreso_editor_proyeccion",
            )
            guardar = st.form_submit_button(
                "Guardar cambios", type="primary", disabled=not actor_nombre
            )
        if guardar:
            originales = {item["id"]: item for item in activos}
            try:
                cambios_totales = 0
                with get_connection() as con:
                    for registro_id, fila in editada.iterrows():
                        registro_id = int(registro_id)
                        original = originales[registro_id]
                        valores_editados = {
                            "confirmados": _entero_no_negativo(
                                fila["Confirmados"], "Confirmados"
                            ),
                            "contacto_nombre": _texto_editor(fila["Nombre de contacto"]),
                            "contacto_celular": _texto_editor(fila["Celular de contacto"]),
                            "observaciones": _texto_editor(fila["Observaciones"]),
                        }
                        cambios = {}
                        if valores_editados["confirmados"] != original["confirmados"]:
                            cambios["confirmados"] = valores_editados["confirmados"]
                        for campo in (
                            "contacto_nombre", "contacto_celular", "observaciones"
                        ):
                            if valores_editados[campo] != (original.get(campo) or ""):
                                cambios[campo] = valores_editados[campo]
                        if not cambios:
                            continue
                        cambios_totales += actualizar_confirmacion_proyeccion(
                            con, registro_id, cambios, actor_id, actor_nombre
                        )
                if cambios_totales:
                    st.success(f"Se guardaron {cambios_totales} cambios en el historial.")
                    st.rerun()
                else:
                    st.info("No se detectaron cambios.")
            except (ValueError, PermissionError) as exc:
                st.error(str(exc))
            except Exception as exc:
                st.error(f"No se pudieron guardar los cambios: {exc}")
    else:
        st.info("No hay instituciones activas. Puedes agregar o reactivar una institución.")

    st.divider()
    with st.expander("Agregar institución"):
        with st.form("form_agregar_institucion_proyeccion", clear_on_submit=True):
            nueva_institucion = st.text_input("Universidad o institución")
            nueva_proyeccion = st.number_input(
                "Estudiantes proyectados", min_value=0, value=0, step=1
            )
            agregar = st.form_submit_button(
                "Agregar institución", type="primary", disabled=not actor_nombre
            )
        if agregar:
            try:
                proyeccion = _entero_no_negativo(
                    nueva_proyeccion, "La proyección"
                )
                with get_connection() as con:
                    crear_institucion_proyeccion(
                        con,
                        nueva_institucion,
                        _normalizar(nueva_institucion),
                        proyeccion,
                        actor_id,
                        actor_nombre,
                    )
                st.success("Institución agregada.")
                st.rerun()
            except (ValueError, PermissionError) as exc:
                st.error(str(exc))
            except Exception as exc:
                st.error(f"No se pudo agregar la institución: {exc}")

    if todos:
        por_id = {item["id"]: item for item in todos}
        with st.expander("Administrar instituciones"):
            administrar_id = st.selectbox(
                "Institución que deseas administrar",
                list(por_id),
                format_func=lambda valor: (
                    f"{por_id[valor]['institucion']}"
                    + (" (inactiva)" if not por_id[valor]["activo"] else "")
                ),
                key="congreso_proyeccion_administrar_id",
            )
            actual = por_id[administrar_id]
            with st.form(f"form_administrar_proyeccion_{administrar_id}"):
                nombre_editar = st.text_input(
                    "Universidad o institución", value=actual["institucion"]
                )
                proyeccion_editar = st.number_input(
                    "Estudiantes proyectados",
                    min_value=0,
                    value=int(actual["proyeccion"]),
                    step=1,
                )
                activo_editar = st.checkbox("Institución activa", value=actual["activo"])
                guardar_admin = st.form_submit_button(
                    "Guardar institución", type="primary", disabled=not actor_nombre
                )
            if guardar_admin:
                try:
                    with get_connection() as con:
                        cantidad = actualizar_institucion_proyeccion(
                            con,
                            administrar_id,
                            nombre_editar,
                            _normalizar(nombre_editar),
                            _entero_no_negativo(proyeccion_editar, "La proyección"),
                            activo_editar,
                            actor_id,
                            actor_nombre,
                        )
                    if cantidad:
                        st.success("Institución actualizada.")
                        st.rerun()
                    else:
                        st.info("No se detectaron cambios.")
                except (ValueError, PermissionError) as exc:
                    st.error(str(exc))
                except Exception as exc:
                    st.error(f"No se pudo actualizar la institución: {exc}")

            inactivos = [item for item in todos if not item["activo"]]
            if inactivos:
                st.markdown("**Instituciones inactivas**")
                st.dataframe(
                    pd.DataFrame(
                        [
                            {
                                "Institución": item["institucion"],
                                "Proyección": item["proyeccion"],
                                "Confirmados": item["confirmados"],
                                "Última actualización": _fecha_hora_ecuador(
                                    item.get("fecha_actualizacion")
                                ),
                            }
                            for item in inactivos
                        ]
                    ),
                    hide_index=True,
                    use_container_width=True,
                )

        if activos:
            activos_por_id = {item["id"]: item for item in activos}
            with st.expander("Eliminar fila de la proyección"):
                st.caption(
                    "La fila dejará de mostrarse, pero conservará su historial y podrá "
                    "reactivarse desde Administrar instituciones."
                )
                with st.form("form_eliminar_fila_proyeccion"):
                    eliminar_id = st.selectbox(
                        "Institución que deseas retirar",
                        list(activos_por_id),
                        format_func=lambda valor: activos_por_id[valor]["institucion"],
                    )
                    confirmar_eliminacion = st.checkbox(
                        "Confirmo que deseo eliminar esta fila de la proyección"
                    )
                    eliminar_fila = st.form_submit_button(
                        "Eliminar fila",
                        disabled=not actor_nombre or not confirmar_eliminacion,
                    )
                if eliminar_fila:
                    fila = activos_por_id[eliminar_id]
                    try:
                        with get_connection() as con:
                            actualizar_institucion_proyeccion(
                                con,
                                eliminar_id,
                                fila["institucion"],
                                fila["institucion_normalizada"],
                                int(fila["proyeccion"]),
                                False,
                                actor_id,
                                actor_nombre,
                            )
                        st.success("Fila eliminada de la proyección.")
                        st.rerun()
                    except (ValueError, PermissionError) as exc:
                        st.error(str(exc))
                    except Exception as exc:
                        st.error(f"No se pudo eliminar la fila: {exc}")


def _vista_historial(oficina_id: str, es_master: bool) -> None:
    with get_connection() as con:
        historial = [dict(item) for item in listar_historial_congreso(
            con, None if es_master else oficina_id
        )]
        invitados = [dict(item) for item in listar_invitados_congreso(
            con, None if es_master else oficina_id
        )]
    if not historial:
        st.info("Aún no hay cambios registrados.")
        return

    tipo_historial = st.selectbox(
        "Mostrar historial",
        ["Todos los cambios", "Proyección estudiantes", "Confirmados estudiantes"],
        key="congreso_tipo_historial",
    )
    historial_mostrado = historial
    if tipo_historial == "Proyección estudiantes":
        historial_mostrado = [
            item for item in historial
            if item.get("entidad_tipo") == "proyeccion_estudiantes"
        ]
    elif tipo_historial == "Confirmados estudiantes":
        historial_mostrado = [
            item for item in historial
            if item.get("entidad_tipo") == "proyeccion_estudiantes"
            and item.get("campo") == "confirmados"
        ]

    if not historial_mostrado:
        st.info("No hay cambios registrados para este filtro.")
        return

    tabla = pd.DataFrame([
        {
            "Fecha (Ecuador)": _fecha_hora_ecuador(item["fecha_cambio"]),
            "Institución": item.get("institucion") or "Directorio",
            "Oficina": _nombre_oficina(item.get("oficina")), "Actualizado por": item["actor_nombre"],
            "Acción": item["accion"],
            "Campo": _ETIQUETAS_CAMPOS.get(item.get("campo"), item.get("campo") or "—"),
            "Valor anterior": item.get("valor_anterior") or "—",
            "Valor nuevo": item.get("valor_nuevo") or "—",
        }
        for item in historial_mostrado
    ])
    st.caption("Fechas y horas mostradas en America/Guayaquil (UTC−5).")
    st.dataframe(tabla, hide_index=True, use_container_width=True, height=480)
    st.download_button(
        "Descargar seguimiento en Excel", data=_datos_exportacion(invitados, historial),
        file_name=(
            f"seguimiento_congreso_"
            f"{datetime.now(ZONA_HORARIA_ECUADOR):%Y%m%d_%H%M}.xlsx"
        ),
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _vista_carga_inicial(actor_nombre: str) -> None:
    with get_connection() as con:
        total_existente = contar_invitados_congreso(con)
    if total_existente:
        st.success(f"La carga inicial ya fue completada. Existen {total_existente} invitados.")
        st.caption("Las actualizaciones posteriores deben realizarse desde Seguimiento.")
        return

    st.markdown("Carga el archivo maestro y revisa el resultado antes de confirmar.")
    archivo = st.file_uploader("Archivo Excel", type=["xlsx"], key="congreso_excel_inicial")
    if archivo is None:
        return
    contenido = archivo.getvalue()
    try:
        resultado = analizar_excel_congreso(contenido)
    except Exception as exc:
        st.error(f"No se pudo leer el archivo: {exc}")
        return

    distribucion = resultado["distribucion"]
    cols = st.columns(5)
    resumen = [
        ("Guayaquil", distribucion.get("guayaquil", 0)),
        ("Portoviejo", distribucion.get("manabi", 0)),
        ("Cuenca", distribucion.get("cuenca", 0)),
        ("Loja", distribucion.get("loja", 0)),
        ("Sin asignar", distribucion.get("sin_asignar", 0)),
    ]
    for col, (etiqueta, cantidad) in zip(cols, resumen):
        col.metric(etiqueta, cantidad)
    st.caption(
        f"{len(resultado['invitados'])} invitados y "
        f"{len(resultado['responsables'])} responsables detectados."
    )
    if resultado["errores"]:
        st.error("Corrige estos errores antes de importar:")
        for error in resultado["errores"]:
            st.write(f"- {error}")
    if resultado["duplicados"]:
        st.error("Se detectaron posibles duplicados:")
        for duplicado in resultado["duplicados"]:
            st.write(f"- {duplicado}")

    puede_importar = not resultado["errores"] and not resultado["duplicados"]
    if st.button(
        "Confirmar carga inicial", type="primary", disabled=not puede_importar,
        key="congreso_confirmar_importacion",
    ):
        try:
            with get_connection() as con:
                cantidad = importar_datos_congreso(
                    con, [dict(item) for item in resultado["invitados"]],
                    resultado["responsables"], archivo.name,
                    sha256(contenido).hexdigest(),
                    actor_nombre or "Carga inicial — Guayaquil",
                )
            st.success(f"Carga completada: {cantidad} invitados importados.")
            st.rerun()
        except ValueError as exc:
            st.error(str(exc))
        except Exception as exc:
            st.error(f"La carga no se completó y no se guardaron datos: {exc}")


def mostrar_seguimiento_congreso() -> None:
    oficina_id = st.session_state.get("oficina_id", "")
    oficina_nombre = st.session_state.get("oficina_nombre", oficina_id)
    es_master = st.session_state.get("oficina_rol") == "master"

    st.markdown(
        """
        <style>
        .congreso-cabecera {border-left: 5px solid #C8A951; padding: .15rem 0 .15rem 1rem; margin-bottom: 1rem;}
        .congreso-cabecera h2 {color: #1A3A5C; margin: 0;}
        .congreso-cabecera p {margin: .25rem 0 0; color: #526170;}
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.markdown(
        f"""
        <div class="congreso-cabecera">
          <h2>Seguimiento Congreso</h2>
          <p>Confirmación de asistentes · 21 y 22 de octubre de 2026 · {oficina_nombre}</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    actor_id, actor_nombre = _actor_actual(oficina_id)
    etiquetas = ["Seguimiento", "Responsables", "Historial"]
    if es_master:
        etiquetas.extend(["Proyección estudiantes", "Carga inicial"])
    pestanas = st.tabs(etiquetas)
    with pestanas[0]:
        _vista_seguimiento(oficina_id, es_master, actor_id, actor_nombre)
    with pestanas[1]:
        _vista_responsables(oficina_id, es_master, actor_id, actor_nombre)
    with pestanas[2]:
        _vista_historial(oficina_id, es_master)
    if es_master:
        with pestanas[3]:
            _vista_proyeccion_estudiantes(actor_id, actor_nombre)
        with pestanas[4]:
            _vista_carga_inicial(actor_nombre)
